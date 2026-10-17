import os
from pathlib import Path
from dotenv import load_dotenv

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, Response
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from datetime import datetime, timedelta
import qrcode
import uuid
import json
import threading
import requests
import queue
import re

try:
    from groq import Groq
    _groq_key = os.getenv("GROQ_API_KEY", "").strip()
    groq_client = Groq(api_key=_groq_key) if _groq_key else None
except Exception as _ge:
    groq_client = None
    print("Groq initialization notice:", _ge)

QR_FOLDER = BASE_DIR / "static" / "qrcodes"
QR_FOLDER.mkdir(parents=True, exist_ok=True)
ISSUES_FOLDER = BASE_DIR / "static" / "issues"
ISSUES_FOLDER.mkdir(parents=True, exist_ok=True)
GOOGLE_SHEET_CONFIG_FILE = BASE_DIR / "google_sheet_config.json"
ONLINE_TIMEOUT_SECONDS = 12


def get_sheet_config():
    if GOOGLE_SHEET_CONFIG_FILE.exists():
        try:
            with open(GOOGLE_SHEET_CONFIG_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {
        "spreadsheet_id": "13Wt2q2Ac76ZWTwUCKxL13dxLyJ6w4ufNebfBsdtjK2E",
        "spreadsheet_url": "https://docs.google.com/spreadsheets/d/13Wt2q2Ac76ZWTwUCKxL13dxLyJ6w4ufNebfBsdtjK2E/edit?gid=0#gid=0",
        "webhook_url": "",
        "last_synced": None
    }


def save_sheet_config(conf):
    try:
        with open(GOOGLE_SHEET_CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(conf, f, indent=2)
    except Exception as e:
        print("Error saving sheet config:", e)


app = Flask(__name__)
app.config["SECRET_KEY"] = "warehouse-ai-v3-change-this"
app.config["SQLALCHEMY_DATABASE_URI"] = f"sqlite:///{BASE_DIR / 'warehouse.db'}"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
db = SQLAlchemy(app)


class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.String(50), unique=True, nullable=False)
    name = db.Column(db.String(100), nullable=False)
    password = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(30), nullable=False)
    qr_token = db.Column(db.String(100), unique=True, nullable=False)
    active = db.Column(db.Boolean, default=True)
    last_seen = db.Column(db.DateTime, nullable=True)

    @property
    def is_online(self):
        if not self.active or not self.last_seen:
            return False
        return (datetime.now() - self.last_seen).total_seconds() < ONLINE_TIMEOUT_SECONDS


class Movement(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    item = db.Column(db.String(150), nullable=False)
    quantity = db.Column(db.Integer, nullable=False)
    direction = db.Column(db.String(10), nullable=False)
    note = db.Column(db.String(500), default="")
    created_by = db.Column(db.String(50), nullable=False)
    status = db.Column(db.String(20), default="Pending")
    verified_by = db.Column(db.String(50), default="")
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class IssueComment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    issue_id = db.Column(db.Integer, db.ForeignKey('issue.id'), nullable=False)
    sender_id = db.Column(db.String(50), nullable=False)
    body = db.Column(db.Text, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class Issue(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    item = db.Column(db.String(150), nullable=False)
    issue_type = db.Column(db.String(50), nullable=False)
    description = db.Column(db.String(500), default="")
    photo = db.Column(db.String(255), default="")
    created_by = db.Column(db.String(50), nullable=False)
    status = db.Column(db.String(30), default="Open")
    resolved_by = db.Column(db.String(50), default="")
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    comments = db.relationship('IssueComment', backref='issue', lazy=True, order_by='IssueComment.created_at.asc()')


class Message(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    sender_id = db.Column(db.String(50), nullable=False)
    receiver_id = db.Column(db.String(50), nullable=False)
    subject = db.Column(db.String(150), nullable=False)
    body = db.Column(db.String(2000), nullable=False)
    is_read = db.Column(db.Boolean, default=False)
    msg_type = db.Column(db.String(20), default="confidential")
    attachment = db.Column(db.String(255), default="")
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    deleted_by_sender = db.Column(db.Boolean, default=False)
    deleted_by_receiver = db.Column(db.Boolean, default=False)


def qr_file(user):
    return QR_FOLDER / f"{user.user_id}.png"


def generate_qr(user):
    data = f"http://127.0.0.1:5000/qr/{user.qr_token}"
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_L,
        box_size=10,
        border=4,
    )
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white").convert('RGB')
    img.save(qr_file(user))


def init_db():
    with app.app_context():
        db.create_all()
        supervisor = User.query.filter_by(user_id="SUP001").first()
        if not supervisor:
            supervisor = User(
                user_id="SUP001",
                name="Main Supervisor",
                password=generate_password_hash("admin123"),
                role="supervisor",
                qr_token=str(uuid.uuid4()),
                active=True,
            )
            db.session.add(supervisor)
            db.session.commit()

        for user in User.query.all():
            if not qr_file(user).exists():
                generate_qr(user)


def me():
    if "user_id" not in session:
        return None
    return db.session.get(User, session["user_id"])


def login_required():
    return me()


def role_dashboard(user):
    if user.role == "supervisor":
        return redirect(url_for("supervisor_dashboard"))
    if user.role == "manager":
        return redirect(url_for("manager_dashboard"))
    return redirect(url_for("employee_dashboard"))


def is_ajax_request():
    return (
        request.headers.get("X-Requested-With") == "XMLHttpRequest"
        or request.is_json
        or request.form.get("is_ajax") == "1"
        or request.args.get("is_ajax") == "1"
        or "application/json" in request.headers.get("Accept", "")
    )


@app.before_request
def update_user_presence():
    uid = session.get("user_id")
    if uid and not request.path.startswith("/static"):
        u = db.session.get(User, uid)
        if u:
            now = datetime.now()
            if not u.last_seen or (now - u.last_seen).total_seconds() > 2:
                u.last_seen = now
                db.session.commit()


@app.context_processor
def inject_global_template_vars():
    try:
        conf = get_sheet_config()
    except Exception:
        conf = {}
    return {
        "sheet_config": conf,
    }


@app.route("/", methods=["GET", "POST"])
def login():
    if session.get("user_id"):
        return role_dashboard(me())

    if request.method == "POST":
        uid = request.form.get("user_id", "").strip().upper()
        password = request.form.get("password", "")
        user = User.query.filter_by(user_id=uid).first()
        if not user and uid == "MGR001":
            user = User.query.filter_by(user_id="MAN001").first()
        elif not user and uid == "MAN001":
            user = User.query.filter_by(user_id="MGR001").first()

        def check_auth(u, p):
            if not u or not u.active:
                return False
            if check_password_hash(u.password, p):
                return True
            # Allow interchangeable default credentials (123 or admin123)
            if p == "123" and check_password_hash(u.password, "admin123"):
                return True
            if p == "admin123" and check_password_hash(u.password, "123"):
                return True
            return False

        if user and check_auth(user, password):
            session["user_id"] = user.id
            user.last_seen = datetime.now()
            db.session.commit()
            trigger_async_sync()
            return role_dashboard(user)

        flash("Invalid User ID or Password.", "error")

    return render_template("login.html")


@app.route("/logout")
def logout():
    uid = session.get("user_id")
    if uid:
        u = db.session.get(User, uid)
        if u:
            u.last_seen = None
            db.session.commit()
            trigger_async_sync()
    session.clear()
    return redirect(url_for("login"))


def dashboard_stats():
    total_in = db.session.query(
        db.func.coalesce(db.func.sum(Movement.quantity), 0)
    ).filter(Movement.direction == "IN", Movement.status != "Deleted").scalar()

    total_out = db.session.query(
        db.func.coalesce(db.func.sum(Movement.quantity), 0)
    ).filter(Movement.direction == "OUT", Movement.status != "Deleted").scalar()

    open_issues = Issue.query.filter_by(status="Open").count()
    online_cutoff = datetime.now() - timedelta(seconds=ONLINE_TIMEOUT_SECONDS)
    employees = User.query.filter(User.role == "employee", User.active == True, User.last_seen >= online_cutoff).count()
    managers = User.query.filter(User.role == "manager", User.active == True, User.last_seen >= online_cutoff).count()

    return int(total_in or 0), int(total_out or 0), open_issues, employees, managers


def get_manager_performance_data():
    manager_users = User.query.filter_by(role="manager").order_by(User.name).all()
    records = []
    for mgr in manager_users:
        verifications = Movement.query.filter_by(verified_by=mgr.user_id, status="Verified").count()
        resolutions = Issue.query.filter_by(resolved_by=mgr.user_id).count()
        flags_raised = Movement.query.filter_by(verified_by=mgr.user_id, status="Flagged").count()
        score = (verifications * 2) + (resolutions * 5)
        work_status = "Online (Working)" if mgr.is_online else "Logged Out"
        account_status = "Active" if mgr.active else "Disabled"
        last_seen = mgr.last_seen.strftime("%Y-%m-%d %H:%M:%S") if mgr.last_seen else "Never"
        records.append({
            "name": mgr.name,
            "user_id": mgr.user_id,
            "work_status": work_status,
            "account_status": account_status,
            "verifications": verifications,
            "resolutions": resolutions,
            "flags_raised": flags_raised,
            "score": score,
            "last_seen": last_seen
        })
    records.sort(key=lambda x: x["score"], reverse=True)
    return records


def get_employee_performance_data():
    employee_users = User.query.filter_by(role="employee").order_by(User.name).all()
    records = []
    for emp in employee_users:
        total_logs = Movement.query.filter(Movement.created_by == emp.user_id, Movement.status != "Deleted").count()
        verified_logs = Movement.query.filter_by(created_by=emp.user_id, status="Verified").count()
        flagged_logs = Movement.query.filter_by(created_by=emp.user_id, status="Flagged").count()
        pending_logs = Movement.query.filter_by(created_by=emp.user_id, status="Pending").count()
        total_units = db.session.query(db.func.coalesce(db.func.sum(Movement.quantity), 0)).filter(
            Movement.created_by == emp.user_id, Movement.status != "Deleted"
        ).scalar() or 0
        reviewed = verified_logs + flagged_logs
        accuracy_rate = f"{round((verified_logs / reviewed) * 100, 1)}%" if reviewed > 0 else ("100%" if verified_logs > 0 else "N/A")
        score = max(0, (verified_logs * 4) + (total_logs * 1) - (flagged_logs * 3) + (int(total_units) // 10))
        work_status = "Online (Working)" if emp.is_online else "Logged Out"
        account_status = "Active" if emp.active else "Disabled"
        last_seen = emp.last_seen.strftime("%Y-%m-%d %H:%M:%S") if emp.last_seen else "Never"
        records.append({
            "name": emp.name,
            "user_id": emp.user_id,
            "work_status": work_status,
            "account_status": account_status,
            "total_logs": total_logs,
            "verified_logs": verified_logs,
            "flagged_logs": flagged_logs,
            "pending_logs": pending_logs,
            "total_units": int(total_units),
            "accuracy_rate": accuracy_rate,
            "score": score,
            "last_seen": last_seen
        })
    records.sort(key=lambda x: x["score"], reverse=True)
    return records


def sync_all_to_google_sheet(webhook_url=None):
    conf = get_sheet_config()
    url = (webhook_url or conf.get("webhook_url", "")).strip()
    if not url:
        return False, "Google Sheet Webhook URL not set. Please deploy the Apps Script in your Google Sheet and paste its URL."

    movements = Movement.query.order_by(Movement.created_at.desc()).all()
    issues = Issue.query.order_by(Issue.created_at.desc()).all()
    users = User.query.order_by(User.name).all()

    # Performance Data
    manager_perf = get_manager_performance_data()
    employee_perf = get_employee_performance_data()

    manager_perf_rows = [
        [
            r["name"],
            r["user_id"],
            r["work_status"],
            r["account_status"],
            r["verifications"],
            r["resolutions"],
            r["flags_raised"],
            r["score"],
            r["last_seen"]
        ]
        for r in manager_perf
    ]

    employee_perf_rows = [
        [
            r["name"],
            r["user_id"],
            r["work_status"],
            r["account_status"],
            r["total_logs"],
            r["verified_logs"],
            r["flagged_logs"],
            r["pending_logs"],
            r["total_units"],
            r["accuracy_rate"],
            r["score"],
            r["last_seen"]
        ]
        for r in employee_perf
    ]

    # 1. MANAGER REPORT
    manager_rows = [
        [
            r["name"],
            r["user_id"],
            r["work_status"],
            r["account_status"],
            r["verifications"],
            r["resolutions"],
            r["score"],
            r["last_seen"]
        ]
        for r in manager_perf
    ]

    # 2. EMPLOYEE REPORT
    employee_users = User.query.filter_by(role="employee").order_by(User.name).all()
    employee_rows = []
    for emp in employee_users:
        emp_movements = Movement.query.filter_by(created_by=emp.user_id).order_by(Movement.created_at.desc()).all()
        work_status = "Online (Working)" if emp.is_online else "Logged Out"
        if emp_movements:
            for m in emp_movements:
                employee_rows.append([
                    emp.name,
                    emp.user_id,
                    work_status,
                    m.created_at.strftime("%Y-%m-%d %H:%M:%S"),
                    m.item,
                    m.quantity,
                    m.direction,
                    m.status,
                    m.verified_by or "Pending Verification",
                    m.note or "-"
                ])
        else:
            employee_rows.append([
                emp.name,
                emp.user_id,
                work_status,
                "-",
                "No Movements Logged Yet",
                0,
                "-",
                "-",
                "-",
                "-"
            ])

    payload = {
        "spreadsheet_id": conf.get("spreadsheet_id", "13Wt2q2Ac76ZWTwUCKxL13dxLyJ6w4ufNebfBsdtjK2E"),
        "action": "full_sync",
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "manager_performance": manager_perf_rows,
        "employee_performance": employee_perf_rows,
        "manager_report": manager_rows,
        "employee_report": employee_rows,
        "movements": [
            [
                m.id,
                m.created_at.strftime("%Y-%m-%d %H:%M:%S"),
                "Movement",
                m.item,
                m.quantity,
                m.direction,
                m.status,
                m.created_by,
                m.verified_by or "-",
                m.note or "-"
            ]
            for m in movements
        ],
        "issues": [
            [
                i.id,
                i.created_at.strftime("%Y-%m-%d %H:%M:%S"),
                "Issue",
                i.item,
                i.issue_type,
                "-",
                i.status,
                i.created_by,
                i.resolved_by or "-",
                i.description or "-"
            ]
            for i in issues
        ],
        "users": [
            [
                u.id,
                u.user_id,
                u.name,
                u.role,
                "Online (Working)" if u.is_online else "Logged Out",
                u.last_seen.strftime("%Y-%m-%d %H:%M:%S") if u.last_seen else "Never"
            ]
            for u in users
        ]
    }

    try:
        resp = requests.post(url, json=payload, timeout=25)
        if resp.status_code in (200, 302):
            conf["last_synced"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            if webhook_url:
                conf["webhook_url"] = webhook_url
            save_sheet_config(conf)
            return True, f"Successfully uploaded {len(movements)} movements, {len(issues)} issues, {len(manager_perf_rows)} manager performance, {len(employee_perf_rows)} employee performance records, and reports to your Google Sheet!"
        return False, f"Google Sheet Webhook returned HTTP {resp.status_code}: {resp.text[:120]}"
    except Exception as e:
        return False, f"Failed to connect to Google Sheet Webhook: {str(e)}"


# ----------------------------------------------------
# REAL-TIME GOOGLE SHEET ENGINE (DEDICATED QUEUE + BATCHING + SMART DEBOUNCE)
# ----------------------------------------------------
sheet_sync_queue = queue.Queue()
sheet_worker_thread = None
sheet_worker_lock = threading.Lock()
_debounce_timer = None
_debounce_lock = threading.Lock()

def _sheet_sync_worker():
    while True:
        try:
            task = sheet_sync_queue.get()
            if task is None:
                break
            
            task_type = task.get("type")
            conf = get_sheet_config()
            url = conf.get("webhook_url", "").strip()

            if url:
                if task_type == "append_rows":
                    rows = task.get("rows", [])
                    event_type = task.get("event_type", "movement")
                    if rows:
                        try:
                            resp = requests.post(url, json={
                                "action": "append_rows",
                                "event_type": event_type,
                                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                                "rows": rows
                            }, timeout=15)
                            if resp.status_code == 200 and resp.json().get("status") == "unknown_action":
                                for r in rows:
                                    requests.post(url, json={
                                        "action": "append_row",
                                        "event_type": event_type,
                                        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                                        "data": r
                                    }, timeout=8)
                        except Exception as e:
                            print("Google Sheet append_rows error:", e)

                elif task_type == "update_status":
                    try:
                        resp = requests.post(url, json={
                            "action": "update_movement_status",
                            "id": task.get("id"),
                            "status": task.get("status"),
                            "verified_by": task.get("verified_by"),
                            "note": task.get("note")
                        }, timeout=10)
                        if resp.status_code == 200 and resp.json().get("status") == "unknown_action":
                            fallback_row = task.get("fallback_row")
                            if fallback_row:
                                requests.post(url, json={
                                    "action": "append_row",
                                    "event_type": "movement",
                                    "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                                    "data": fallback_row
                                }, timeout=8)
                    except Exception as e:
                        print("Google Sheet update_status error:", e)

                elif task_type == "append_single":
                    try:
                        requests.post(url, json={
                            "action": "append_row",
                            "event_type": task.get("event_type", "movement"),
                            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                            "data": task.get("data")
                        }, timeout=10)
                    except Exception as e:
                        print("Google Sheet append_single error:", e)

                elif task_type == "full_sync":
                    try:
                        with app.app_context():
                            sync_all_to_google_sheet()
                    except Exception as e:
                        print("Google Sheet full_sync error:", e)

            sheet_sync_queue.task_done()
        except Exception as e:
            print("Google Sheet sync worker loop error:", e)


def ensure_sheet_worker():
    global sheet_worker_thread
    with sheet_worker_lock:
        if sheet_worker_thread is None or not sheet_worker_thread.is_alive():
            sheet_worker_thread = threading.Thread(target=_sheet_sync_worker, daemon=True)
            sheet_worker_thread.start()


ensure_sheet_worker()


def push_live_event_to_google_sheet(event_type, details):
    ensure_sheet_worker()
    sheet_sync_queue.put({
        "type": "append_single",
        "event_type": event_type,
        "data": details
    })


def push_live_movements_to_google_sheet(movements):
    ensure_sheet_worker()
    if not movements:
        return
    rows = [
        [
            m.id,
            m.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            "Movement",
            m.item,
            m.quantity,
            m.direction,
            m.status,
            m.created_by,
            m.verified_by or "-",
            m.note or "-"
        ]
        for m in movements
    ]
    sheet_sync_queue.put({
        "type": "append_rows",
        "event_type": "movement",
        "rows": rows
    })


def update_google_sheet_movement_status(movement_id, status, verified_by="", note="", fallback_row=None):
    ensure_sheet_worker()
    sheet_sync_queue.put({
        "type": "update_status",
        "id": movement_id,
        "status": status,
        "verified_by": verified_by,
        "note": note,
        "fallback_row": fallback_row
    })


def trigger_async_sync(delay_seconds=3.0):
    """
    Debounced full sync of background reports and performance sheets.
    Prevents multiple overlapping heavy syncs from stalling live event delivery.
    """
    global _debounce_timer
    ensure_sheet_worker()

    def _enqueue_full_sync():
        sheet_sync_queue.put({"type": "full_sync"})

    with _debounce_lock:
        if _debounce_timer and _debounce_timer.is_alive():
            _debounce_timer.cancel()
        _debounce_timer = threading.Timer(delay_seconds, _enqueue_full_sync)
        _debounce_timer.daemon = True
        _debounce_timer.start()


def start_google_sheet_autosync():
    def _loop():
        import time
        while True:
            try:
                time.sleep(30)
                conf = get_sheet_config()
                if conf.get("webhook_url"):
                    ensure_sheet_worker()
                    sheet_sync_queue.put({"type": "full_sync"})
            except Exception as e:
                print("Auto-sync loop error:", e)
    threading.Thread(target=_loop, daemon=True).start()


# ----------------------------------------------------
# REAL-TIME LIVE ACTIVITY STREAM (SSE + WEBSOCKET-LIKE PUSH)
# ----------------------------------------------------
activity_subscribers = []
activity_subscribers_lock = threading.Lock()


def notify_activity_subscribers(event_data):
    with activity_subscribers_lock:
        dead_queues = []
        for q in activity_subscribers:
            try:
                q.put_nowait(event_data)
            except Exception:
                dead_queues.append(q)
        for dq in dead_queues:
            if dq in activity_subscribers:
                activity_subscribers.remove(dq)


@app.route("/api/live_activity/stream")
def live_activity_stream():
    user = login_required()
    if not user:
        return jsonify({"error": "Unauthorized"}), 401

    def event_stream():
        q = queue.Queue(maxsize=100)
        with activity_subscribers_lock:
            activity_subscribers.append(q)
        try:
            # Instruct browser EventSource to reconnect within 1000ms if dropped
            yield "retry: 1000\n\n"
            # Yield initial connection confirmation
            yield f"event: connected\ndata: {json.dumps({'status': 'connected', 'user': user.user_id, 'time': datetime.now().strftime('%H:%M:%S')})}\n\n"
            while True:
                try:
                    event = q.get(timeout=10)
                    yield f"event: activity\ndata: {json.dumps(event)}\n\n"
                except queue.Empty:
                    # Heartbeat comment to keep HTTP connection alive
                    yield ": ping\n\n"
        except GeneratorExit:
            pass
        finally:
            with activity_subscribers_lock:
                if q in activity_subscribers:
                    activity_subscribers.remove(q)

    return Response(event_stream(), mimetype="text/event-stream", headers={
        "Cache-Control": "no-cache, no-transform",
        "X-Accel-Buffering": "no",
        "Connection": "keep-alive"
    })


@app.route("/dashboard")
def dashboard():
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    return role_dashboard(user)


@app.route("/export/weekly", methods=["POST"])
def export_weekly():
    user = login_required()
    if not user or user.role != "manager":
        return redirect(url_for("dashboard"))
    
    import csv, time
    from pathlib import Path
    
    movements = Movement.query.order_by(Movement.created_at.desc()).all()
    issues = Issue.query.order_by(Issue.created_at.desc()).all()
    
    filename = f"weekly_report_{int(time.time())}.csv"
    filepath = Path(app.root_path) / "static" / "reports" / filename
    
    with open(filepath, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(["--- LIVE ACTIVITY (MOVEMENTS) ---"])
        writer.writerow(["Date", "Item", "Quantity", "Direction", "Status", "Verified By"])
        for m in movements:
            writer.writerow([m.created_at.strftime("%Y-%m-%d %H:%M"), m.item, m.quantity, m.direction, m.status, m.verified_by])
            
        writer.writerow([])
        writer.writerow(["--- ISSUES & FLAGS ---"])
        writer.writerow(["Date", "Item", "Type", "Description", "Status", "Resolved By"])
        for i in issues:
            writer.writerow([i.created_at.strftime("%Y-%m-%d %H:%M"), i.item, i.issue_type, i.description, i.status, i.resolved_by])
            
    supervisors = User.query.filter_by(role="supervisor").all()
    for sup in supervisors:
        msg = Message(
            sender_id=user.user_id,
            receiver_id=sup.user_id,
            subject=f"Weekly Report ({datetime.utcnow().strftime('%d %b')})",
            body="Manager submitted the weekly activity and issues report.",
            attachment=filename,
            msg_type="confidential"
        )
        db.session.add(msg)
    db.session.commit()
    
    flash("Weekly report compiled and securely sent to Supervisors.", "success")
    return redirect(url_for("dashboard"))


@app.route("/send_report/employee", methods=["POST"])
def send_employee_report():
    user = login_required()
    if not user or user.role != "manager":
        flash("Unauthorized.", "error")
        return redirect(url_for("login"))
        
    import csv, time
    from pathlib import Path
    
    users = User.query.filter_by(role="employee").order_by(User.name).all()
    filename = f"employee_report_{int(time.time())}.csv"
    filepath = Path(app.root_path) / "static" / "reports" / filename
    
    with open(filepath, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(["Employee Name", "Employee ID", "Action Date", "Item Logged", "Quantity", "Direction", "Verification Status", "Verified By"])
        for u in users:
            movements = Movement.query.filter_by(created_by=u.user_id).order_by(Movement.created_at.desc()).all()
            for m in movements:
                writer.writerow([u.name, u.user_id, m.created_at.strftime("%Y-%m-%d %H:%M"), m.item, m.quantity, m.direction, m.status, m.verified_by])
                
    supervisors = User.query.filter_by(role="supervisor").all()
    for sup in supervisors:
        msg = Message(
            sender_id=user.user_id,
            receiver_id=sup.user_id,
            subject="Employee Report",
            body="Manager submitted the Employee Report.",
            attachment=filename,
            msg_type="confidential"
        )
        db.session.add(msg)
    db.session.commit()
    
    flash("Employee Report compiled and securely sent to Supervisors.", "success")
    return redirect(url_for("dashboard"))


@app.route("/export/users/<role>", methods=["POST"])
def export_users(role):
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    
    if role not in ["employee", "manager"]:
        flash("Invalid report type.", "error")
        return redirect(url_for("dashboard"))
        
    if role == "manager" and user.role != "supervisor":
        flash("Unauthorized to export manager report.", "error")
        return redirect(url_for("dashboard"))
        
    if user.role not in ["supervisor", "manager"]:
        return redirect(url_for("dashboard"))
    
    import csv, time
    from pathlib import Path
    from flask import send_file
    
    users = User.query.filter_by(role=role).order_by(User.name).all()
    filename = f"{role}_report_{int(time.time())}.csv"
    filepath = Path(app.root_path) / "static" / "reports" / filename
    
    with open(filepath, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        if role == "employee":
            writer.writerow(["Employee Name", "Employee ID", "Action Date", "Item Logged", "Quantity", "Direction", "Verification Status", "Verified By"])
        else:
            writer.writerow(["Name", "User ID", "Status", "Logs Verified", "Issues Resolved", "Performance Score"])
        
        for u in users:
            status = "Active" if u.active else "Disabled"
            if role == "employee":
                movements = Movement.query.filter_by(created_by=u.user_id).order_by(Movement.created_at.desc()).all()
                for m in movements:
                    writer.writerow([u.name, u.user_id, m.created_at.strftime("%Y-%m-%d %H:%M"), m.item, m.quantity, m.direction, m.status, m.verified_by])
            else:
                verified = Movement.query.filter_by(verified_by=u.user_id).count()
                resolved = Issue.query.filter_by(resolved_by=u.user_id).count()
                score = (verified * 2) + (resolved * 5)
                writer.writerow([u.name, u.user_id, status, verified, resolved, score])
            
    return send_file(filepath, as_attachment=True)


@app.route("/supervisor")
def supervisor_dashboard():
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    if user.role != "supervisor":
        return role_dashboard(user)

    if request.args.get("msg_type") or request.args.get("receiver_id"):
        return redirect(url_for("supervisor_dashboard"))

    tin, tout, issues_count, employees, managers = dashboard_stats()
    users = User.query.order_by(User.role, User.name).all()
    movements = Movement.query.order_by(Movement.id.desc()).limit(50).all()
    open_issues = Issue.query.filter_by(status="Open").order_by(Issue.id.desc()).all()
    resolved_issues = Issue.query.filter_by(status="Resolved").order_by(Issue.id.desc()).limit(15).all()
    
    # Separate normal messages and report attachments
    all_conf_messages = Message.query.filter(Message.msg_type=="confidential", Message.attachment=="").order_by(Message.id.desc()).all()
    
    unique_conversations = []
    seen_pairs = set()
    for m in all_conf_messages:
        pair = tuple(sorted([m.sender_id, m.receiver_id]))
        if pair not in seen_pairs:
            seen_pairs.add(pair)
            unique_conversations.append(m)
            if len(unique_conversations) >= 20:
                break
                
    messages = unique_conversations
    reports = Message.query.filter(Message.receiver_id==user.user_id, Message.attachment!="").order_by(Message.id.desc()).all()
    manager_reports = [r for r in reports if "Weekly" in r.subject or "Manager" in r.subject]
    employee_reports = [r for r in reports if "Employee" in r.subject]
    
    common_messages = Message.query.filter_by(msg_type="common").order_by(Message.id.desc()).limit(10).all()
    unread = Message.query.filter(Message.receiver_id==user.user_id, Message.msg_type=="confidential", Message.attachment=="", Message.is_read==False).count()
    
    online_cutoff = datetime.now() - timedelta(seconds=ONLINE_TIMEOUT_SECONDS)
    active_users = User.query.filter(User.active == True, User.last_seen >= online_cutoff).order_by(User.name).all()
    offline_users = User.query.filter(User.active == True, db.or_(User.last_seen == None, User.last_seen < online_cutoff)).order_by(User.name).all()

    # Calculate Performance Data
    manager_perf = get_manager_performance_data()
    employee_perf = get_employee_performance_data()

    verified_logs_count = Movement.query.filter_by(status="Verified").count()
    flagged_logs_count = Movement.query.filter_by(status="Flagged").count()
    deleted_logs_count = Movement.query.filter_by(status="Deleted").count()
    pending_logs_count = Movement.query.filter_by(status="Pending").count()
    total_movements_count = Movement.query.filter(Movement.status != "Deleted").count()
    reviewed_count = verified_logs_count + flagged_logs_count
    accuracy_rate = f"{round((verified_logs_count / reviewed_count) * 100, 1)}%" if reviewed_count > 0 else "100.0%"
    verified_pct = f"{round((verified_logs_count / max(1, total_movements_count)) * 100, 1)}%"
    flagged_pct = f"{round((flagged_logs_count / max(1, total_movements_count)) * 100, 1)}%"

    sheet_config = get_sheet_config()

    return render_template(
        "supervisor.html",
        user=user, total_in=tin, total_out=tout, open_issues=issues_count,
        employees=employees, managers=managers, users=users, active_users=active_users, offline_users=offline_users,
        movements=movements, issues_list=open_issues, resolved_issues=resolved_issues,
        messages=messages, manager_reports=manager_reports, employee_reports=employee_reports, common_messages=common_messages, unread=unread,
        manager_perf=manager_perf, employee_perf=employee_perf, sheet_config=sheet_config,
        verified_logs_count=verified_logs_count, flagged_logs_count=flagged_logs_count,
        deleted_logs_count=deleted_logs_count, pending_logs_count=pending_logs_count,
        total_movements_count=total_movements_count, accuracy_rate=accuracy_rate,
        verified_pct=verified_pct, flagged_pct=flagged_pct
    )


@app.route("/manager")
def manager_dashboard():
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    if user.role not in {"manager", "supervisor"}:
        return role_dashboard(user)

    if request.args.get("msg_type") or request.args.get("receiver_id"):
        return redirect(url_for("manager_dashboard"))

    tin, tout, issues, employees, managers = dashboard_stats()
    movements = Movement.query.filter(Movement.status != "Deleted").order_by(Movement.id.desc()).limit(30).all()
    open_issues = Issue.query.filter_by(status="Open").order_by(Issue.id.desc()).all()
    messages = Message.query.filter(
        db.or_(
            db.and_(Message.receiver_id == user.user_id, Message.deleted_by_receiver == False),
            db.and_(Message.sender_id == user.user_id, Message.deleted_by_sender == False)
        ),
        Message.msg_type == "confidential"
    ).order_by(Message.id.desc()).limit(15).all()
    common_messages = Message.query.filter_by(msg_type="common").order_by(Message.id.desc()).limit(10).all()
    unread = Message.query.filter_by(receiver_id=user.user_id, msg_type="confidential", is_read=False).count()
    
    online_cutoff = datetime.now() - timedelta(seconds=ONLINE_TIMEOUT_SECONDS)
    active_users = User.query.filter(User.active == True, User.last_seen >= online_cutoff).order_by(User.name).all()
    offline_users = User.query.filter(User.active == True, db.or_(User.last_seen == None, User.last_seen < online_cutoff)).order_by(User.name).all()

    employee_perf = get_employee_performance_data()
    manager_perf = get_manager_performance_data()

    return render_template(
        "manager.html",
        user=user, total_in=tin, total_out=tout, open_issues=issues,
        employees=employees, managers=managers, active_users=active_users, offline_users=offline_users,
        movements=movements, issues_list=open_issues,
        messages=messages, common_messages=common_messages, unread=unread,
        employee_perf=employee_perf, manager_perf=manager_perf, sheet_config=get_sheet_config()
    )


@app.route("/employee")
def employee_dashboard():
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    if user.role != "employee":
        return role_dashboard(user)

    if request.args.get("msg_type") or request.args.get("receiver_id"):
        return redirect(url_for("employee_dashboard"))

    my_movements = Movement.query.filter(
        Movement.created_by == user.user_id,
        Movement.status != "Deleted"
    ).order_by(Movement.id.desc()).limit(20).all()

    my_issues = Issue.query.filter_by(
        created_by=user.user_id, status="Open"
    ).order_by(Issue.id.desc()).limit(10).all()

    messages = Message.query.filter(
        db.or_(
            db.and_(Message.receiver_id == user.user_id, Message.deleted_by_receiver == False),
            db.and_(Message.sender_id == user.user_id, Message.deleted_by_sender == False)
        ),
        Message.msg_type == "confidential"
    ).order_by(Message.id.desc()).limit(15).all()

    common_messages = Message.query.filter_by(msg_type="common").order_by(Message.id.desc()).limit(10).all()

    unread = Message.query.filter_by(
        receiver_id=user.user_id, msg_type="confidential", is_read=False
    ).count()

    online_cutoff = datetime.now() - timedelta(seconds=ONLINE_TIMEOUT_SECONDS)
    active_users = User.query.filter(User.active == True, User.last_seen >= online_cutoff).order_by(User.name).all()
    offline_users = User.query.filter(User.active == True, db.or_(User.last_seen == None, User.last_seen < online_cutoff)).order_by(User.name).all()

    return render_template(
        "employee.html",
        user=user, movements=my_movements,
        issues=my_issues, messages=messages, 
        common_messages=common_messages, unread=unread,
        active_users=active_users, offline_users=offline_users
    )


@app.route("/api/dashboard")
def dashboard_api():
    user = login_required()
    if not user:
        return jsonify({"logged_in": False}), 401

    tin, tout, issues, employees, managers = dashboard_stats()
    latest = Movement.query.filter(Movement.status != "Deleted").order_by(Movement.id.desc()).first()
    unread = Message.query.filter_by(receiver_id=user.user_id, is_read=False).count()

    online_cutoff = datetime.now() - timedelta(seconds=ONLINE_TIMEOUT_SECONDS)
    online_users = User.query.filter(User.active == True, User.last_seen >= online_cutoff).order_by(User.name).all()
    offline_users = User.query.filter(User.active == True, db.or_(User.last_seen == None, User.last_seen < online_cutoff)).order_by(User.name).all()
    active_user_list = [{"id": u.id, "user_id": u.user_id, "name": u.name, "role": u.role} for u in online_users]
    offline_user_list = [{"id": u.id, "user_id": u.user_id, "name": u.name, "role": u.role} for u in offline_users]

    if user.role == "employee":
        movements_query = Movement.query.filter(
            Movement.status != "Deleted",
            Movement.created_by == user.user_id
        ).order_by(Movement.id.desc()).limit(30).all()
    elif user.role == "supervisor":
        movements_query = Movement.query.order_by(Movement.id.desc()).limit(50).all()
    else:
        movements_query = Movement.query.filter(
            Movement.status != "Deleted"
        ).order_by(Movement.id.desc()).limit(35).all()

    recent_movements = [
        {
            "id": m.id,
            "date": m.created_at.strftime("%d %b %Y"),
            "time": m.created_at.strftime("%H:%M:%S"),
            "short_time": m.created_at.strftime("%d %b %H:%M"),
            "item": m.item,
            "quantity": m.quantity,
            "direction": m.direction,
            "status": m.status,
            "created_by": m.created_by,
            "verified_by": m.verified_by or "",
            "note": m.note or ""
        }
        for m in movements_query
    ]

    common_msgs = Message.query.filter_by(msg_type="common").order_by(Message.id.desc()).limit(15).all()
    common_list = [
        {
            "id": cm.id,
            "sender_id": cm.sender_id,
            "subject": cm.subject,
            "body": cm.body,
            "created_at": cm.created_at.strftime("%d %b %H:%M")
        }
        for cm in common_msgs
    ]

    manager_perf = get_manager_performance_data() if user.role in ["supervisor", "manager"] else []
    employee_perf = get_employee_performance_data() if user.role in ["supervisor", "manager"] else []

    verified_logs_count = Movement.query.filter_by(status="Verified").count()
    flagged_logs_count = Movement.query.filter_by(status="Flagged").count()
    deleted_logs_count = Movement.query.filter_by(status="Deleted").count()
    pending_logs_count = Movement.query.filter_by(status="Pending").count()
    total_movements_count = Movement.query.filter(Movement.status != "Deleted").count()
    reviewed_count = verified_logs_count + flagged_logs_count
    accuracy_rate = f"{round((verified_logs_count / reviewed_count) * 100, 1)}%" if reviewed_count > 0 else "100.0%"
    verified_pct = f"{round((verified_logs_count / max(1, total_movements_count)) * 100, 1)}%"
    flagged_pct = f"{round((flagged_logs_count / max(1, total_movements_count)) * 100, 1)}%"

    return jsonify({
        "logged_in": True,
        "total_in": tin,
        "total_out": tout,
        "open_issues": issues,
        "employees": employees,
        "managers": managers,
        "active_users": active_user_list,
        "active_users_count": len(active_user_list),
        "offline_users": offline_user_list,
        "recent_movements": recent_movements,
        "common_messages": common_list,
        "latest_id": latest.id if latest else 0,
        "latest_item": latest.item if latest else "",
        "latest_quantity": latest.quantity if latest else 0,
        "latest_direction": latest.direction if latest else "",
        "latest_by": latest.created_by if latest else "",
        "unread": unread,
        "manager_perf": manager_perf,
        "employee_perf": employee_perf,
        "verified_logs_count": verified_logs_count,
        "flagged_logs_count": flagged_logs_count,
        "deleted_logs_count": deleted_logs_count,
        "pending_logs_count": pending_logs_count,
        "total_movements_count": total_movements_count,
        "accuracy_rate": accuracy_rate,
        "verified_pct": verified_pct,
        "flagged_pct": flagged_pct
    })


@app.route("/api/presence/ping", methods=["GET", "POST"])
def presence_ping():
    uid = session.get("user_id")
    if not uid:
        return jsonify({"ok": False}), 401
    u = db.session.get(User, uid)
    if u:
        u.last_seen = datetime.now()
        db.session.commit()
        return jsonify({"ok": True, "user_id": u.user_id})
    return jsonify({"ok": False}), 404


@app.route("/api/presence/offline", methods=["GET", "POST"])
def presence_offline():
    uid = session.get("user_id")
    if uid:
        u = db.session.get(User, uid)
        if u:
            u.last_seen = None
            db.session.commit()
    return jsonify({"ok": True})


def extract_movements_from_bytes(file_bytes, filename, default_direction="IN"):
    import re, io, csv
    records = []
    errors = []
    fn = filename.lower()
    raw_rows = []

    if fn.endswith(('.csv', '.tsv', '.txt')):
        try:
            text = file_bytes.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode('utf-8')
            except UnicodeDecodeError:
                text = file_bytes.decode('latin-1', errors='replace')

        sample = text[:4096]
        counts = {',': sample.count(','), ';': sample.count(';'), '\t': sample.count('\t'), '|': sample.count('|')}
        best_delim = max(counts, key=counts.get) if max(counts.values()) > 0 else (',' if not fn.endswith('.tsv') else '\t')
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
            delimiter = dialect.delimiter
        except Exception:
            delimiter = best_delim

        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        raw_rows = [[str(c).strip() for c in row] for row in reader if any(str(c).strip() for c in row)]
    elif fn.endswith(('.xlsx', '.xlsm', '.xltx')):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None and str(cell).strip() != '' for cell in row):
                raw_rows.append([str(c).strip() if c is not None else '' for c in row])
    elif fn.endswith('.xls'):
        import pandas as pd
        df = pd.read_excel(io.BytesIO(file_bytes))
        raw_rows = [[str(c) for c in df.columns]]
        for row in df.values:
            raw_rows.append([str(c) if pd.notna(c) else '' for c in row])
    else:
        return [], ['Unsupported file format. Please upload .xlsx, .xls, or .csv files.']

    if not raw_rows:
        return [], ['The uploaded file is empty.']

    def normalize_h(h):
        return str(h).strip().lower().replace('_', ' ').replace('-', ' ').replace('/', ' ')

    ITEM_ALIASES = ['item', 'product', 'sku', 'material', 'name', 'part', 'description', 'desc', 'title', 'goods', 'article', 'particular', 'model', 'code', 'barcode', 'asset', 'inventory']
    QTY_ALIASES = ['quantity', 'qty', 'count', 'units', 'amount', 'pcs', 'pieces', 'vol', 'volume', 'total', 'nos', 'number', 'boxes', 'carton', 'cartons', 'packets', 'pkg']
    DIR_ALIASES = ['direction', 'type', 'movement', 'in out', 'flow', 'action', 'status', 'io', 'inward', 'outward', 'inbound', 'outbound', 'trans', 'mode']
    NOTE_ALIASES = ['note', 'remark', 'comment', 'details', 'reason', 'batch', 'po', 'location', 'memo', 'supplier', 'dest', 'destination', 'ref', 'reference']

    header_idx = -1
    best_item_col = -1
    best_qty_col = -1
    best_dir_col = -1
    best_note_col = -1

    max_scan = min(len(raw_rows), 10)
    for r_idx in range(max_scan):
        row = [normalize_h(c) for c in raw_rows[r_idx]]
        i_col, q_col, d_col, n_col = -1, -1, -1, -1
        for c_idx, cell in enumerate(row):
            if not cell:
                continue
            if i_col == -1 and any(a in cell for a in ITEM_ALIASES):
                i_col = c_idx
            elif q_col == -1 and any(a in cell for a in QTY_ALIASES):
                q_col = c_idx
            elif d_col == -1 and any(a in cell for a in DIR_ALIASES):
                d_col = c_idx
            elif n_col == -1 and any(a in cell for a in NOTE_ALIASES):
                n_col = c_idx

        if i_col != -1 and q_col != -1:
            header_idx = r_idx
            best_item_col, best_qty_col, best_dir_col, best_note_col = i_col, q_col, d_col, n_col
            break
        elif (i_col != -1 or q_col != -1) and header_idx == -1:
            header_idx = r_idx
            best_item_col, best_qty_col, best_dir_col, best_note_col = i_col, q_col, d_col, n_col

    if header_idx != -1:
        item_col = best_item_col
        qty_col = best_qty_col
        dir_col = best_dir_col
        note_col = best_note_col
        data_rows = raw_rows[header_idx + 1:]
        start_row_num = header_idx + 2
    else:
        item_col = 0
        qty_col = 1 if len(raw_rows[0]) > 1 else -1
        dir_col = 2 if len(raw_rows[0]) > 2 else -1
        note_col = 3 if len(raw_rows[0]) > 3 else -1
        data_rows = raw_rows
        start_row_num = 1

    if item_col == -1: item_col = 0
    if qty_col == -1 and len(raw_rows[0]) > 1: qty_col = 1

    for row_offset, row in enumerate(data_rows):
        row_idx = start_row_num + row_offset
        if not row or all(not str(c).strip() for c in row):
            continue

        item_val = row[item_col].strip() if item_col < len(row) else ''
        if not item_val:
            errors.append(f'Row {row_idx}: Skipped (Item Name is empty)')
            continue

        qty_str = row[qty_col].strip() if (qty_col != -1 and qty_col < len(row)) else '0'
        num_match = re.search(r'[-+]?\d*\.?\d+', qty_str.replace(',', ''))
        if num_match:
            try:
                qty_val = int(round(float(num_match.group(0))))
            except Exception:
                qty_val = 0
        else:
            qty_val = 0

        if qty_val <= 0:
            errors.append(f'Row {row_idx} ({item_val}): Skipped (Invalid quantity "{qty_str}")')
            continue

        dir_val = default_direction
        if dir_col != -1 and dir_col < len(row) and row[dir_col].strip():
            raw_dir = row[dir_col].strip().upper()
            if any(k in raw_dir for k in ['OUT', 'DISP', 'SEND', 'EXP', 'LEAVE', '-', 'OUTWARD', 'OUTBOUND', 'DISPATCH', 'EXIT', 'EXPORT', 'ISSUE', 'DELIVERY', 'SHIPPED']):
                dir_val = 'OUT'
            elif any(k in raw_dir for k in ['IN', 'REC', 'ARR', 'ENT', '+', 'INWARD', 'INBOUND', 'RECEIVE', 'ENTRY', 'RESTOCK', 'PURCHASE', 'IMPORT', 'RETURN']):
                dir_val = 'IN'

        note_val = row[note_col].strip() if note_col != -1 and note_col < len(row) else ''
        records.append({'item': item_val, 'quantity': qty_val, 'direction': dir_val, 'note': note_val})

    return records, errors


def fallback_parse_multi_items(user_msg, default_dir="IN"):
    text = (user_msg or "").strip()
    lower = text.lower()

    in_cues = ["add", "added", "receive", "received", "got", "inbound", "incoming", "arrived", "restock", "inward", "unloaded", "entered"]
    out_cues = ["dispatched", "dispatch", "sent", "send", "outbound", "outgoing", "shipped", "remove", "removed", "took", "delivered", "exit"]

    global_dir = default_dir
    if any(re.search(r'\b' + c + r'\b', lower) for c in out_cues):
        global_dir = "OUT"
    elif any(re.search(r'\b' + c + r'\b', lower) for c in in_cues):
        global_dir = "IN"

    cleaned = re.sub(r'^(?:please\s+)?(?:add|log|record|received?|dispatched?|insert)\s+', '', text, flags=re.IGNORECASE).strip()
    segments = re.split(r'[,;\n]|\band\b', cleaned, flags=re.IGNORECASE)
    items = []
    missing_criteria = []

    for seg in segments:
        s = seg.strip()
        if not s:
            continue

        seg_dir = global_dir
        if any(re.search(r'\b' + c + r'\b', s.lower()) for c in out_cues):
            seg_dir = "OUT"
        elif any(re.search(r'\b' + c + r'\b', s.lower()) for c in in_cues):
            seg_dir = "IN"

        qty_match = re.search(r'\b(\d+)\s*(?:units?|pcs?|pieces?|boxes?|items?|cartons?|nos?|pkgs?)?\b', s, re.IGNORECASE)
        qty = int(qty_match.group(1)) if qty_match else None

        note_match = re.search(r'\b(?:in|at|for|to|from|bay|dock|sector|po|loc|location)\s+([a-zA-Z0-9\s#-]+)', s, re.IGNORECASE)
        note = note_match.group(0).strip() if note_match else ""

        item_text = s
        if qty_match:
            item_text = item_text.replace(qty_match.group(0), '')
        if note_match:
            item_text = item_text.replace(note_match.group(0), '')

        item_text = re.sub(r'\b(?:add|received?|dispatched?|in|out|of|some|units?|pcs?|pieces?|boxes?|the|a|an)\b', '', item_text, flags=re.IGNORECASE)
        item_text = re.sub(r'[^\w\s-]', '', item_text).strip()

        if item_text:
            item_title = item_text.title()
            is_comp = bool(qty and qty > 0)
            items.append({
                "item": item_title,
                "quantity": qty,
                "direction": seg_dir,
                "note": note,
                "is_complete": is_comp
            })
            if not qty:
                missing_criteria.append(f"Quantity for '{item_title}'")

    if not items and text:
        items.append({
            "item": text.title(),
            "quantity": None,
            "direction": global_dir,
            "note": "",
            "is_complete": False
        })
        missing_criteria.append("Quantity")

    return items, missing_criteria


def parse_movement_with_ai(user_msg, history=None, current_draft=None):
    current_draft = dict(current_draft or {})

    if groq_client:
        system_prompt = (
            "You are an intelligent, friendly warehouse inventory movement conversational assistant.\n"
            "The employee may prompt you with a single item or multiple items to add, receive, or dispatch at once "
            "(e.g., 'add 70 sofa, 80 car, 90 wheel' or 'dispatched 15 helmets and received 20 jacks').\n\n"
            "Your job is:\n"
            "1. Extract ALL mentioned items into the 'items' list.\n"
            "   Each element of 'items' must have:\n"
            "   - item: (string, Title Cased name / SKU, e.g. 'Sofa', 'Car', 'Wheel')\n"
            "   - quantity: (integer > 0, or null if unspecified)\n"
            "   - direction: ('IN' for add/received/got/restock/inbound, 'OUT' for dispatched/sent/shipped; default to 'IN' if 'add' or neutral)\n"
            "   - note: (string, any bay, PO, location, or '')\n"
            "   - is_complete: (boolean, true ONLY if item is non-empty and quantity is an integer > 0 and direction in ['IN', 'OUT'])\n"
            "2. Identify criteria that are left/missing in 'missing_criteria':\n"
            "   - Note any required criteria missing (e.g. 'Quantity for Wheels').\n"
            "   - Note any optional criteria missing (e.g. 'Notes / Bay / PO (Optional)').\n"
            "3. Formulate a friendly 'reply_message':\n"
            "   - State clearly which items are being added and auto-confirmed.\n"
            "   - Specifically ask for criteria that are left (e.g. Item Name / SKU, Quantity, Movement Direction, Notes / Bay / PO).\n\n"
            "Respond ONLY with valid JSON with keys: items (array), missing_criteria (array of strings), reply_message (string)."
        )

        messages = [{"role": "system", "content": system_prompt}]

        draft_context = []
        if current_draft.get("item"):
            draft_context.append(f"item='{current_draft['item']}'")
        if current_draft.get("quantity"):
            draft_context.append(f"quantity={current_draft['quantity']}")
        if current_draft.get("direction"):
            draft_context.append(f"direction='{current_draft['direction']}'")
        if current_draft.get("note"):
            draft_context.append(f"note='{current_draft['note']}'")

        if draft_context:
            messages.append({
                "role": "system",
                "content": f"Existing draft state: {', '.join(draft_context)}. Merge new details into this state if applicable."
            })

        if history and isinstance(history, list):
            for h in history[-6:]:
                role = "assistant" if h.get("role") in ("assistant", "ai") else "user"
                content = str(h.get("content", "")).strip()
                if content:
                    messages.append({"role": role, "content": content})

        messages.append({"role": "user", "content": str(user_msg or "").strip()})

        for model_candidate in ["qwen/qwen3.8-27b", "openai/gpt-oss-20b"]:
            try:
                res = groq_client.chat.completions.create(
                    model=model_candidate,
                    messages=messages,
                    temperature=0.1,
                    response_format={"type": "json_object"},
                    timeout=10
                )
                raw_json = res.choices[0].message.content
                parsed = json.loads(raw_json)

                items = parsed.get("items")
                if not items or not isinstance(items, list):
                    if "item" in parsed:
                        items = [parsed]
                    else:
                        items = []

                cleaned_items = []
                for it in items:
                    it_name = str(it.get("item") or "").strip().title()
                    it_qty = it.get("quantity")
                    try:
                        it_qty = int(it_qty) if it_qty is not None else None
                    except Exception:
                        it_qty = None
                    it_dir = str(it.get("direction") or "IN").upper().strip()
                    if it_dir not in ("IN", "OUT"):
                        it_dir = "IN"
                    it_note = str(it.get("note") or "").strip()
                    is_c = bool(it_name and it_qty and it_qty > 0 and it_dir in ("IN", "OUT"))
                    cleaned_items.append({
                        "item": it_name,
                        "quantity": it_qty,
                        "direction": it_dir,
                        "note": it_note,
                        "is_complete": is_c
                    })

                missing_criteria = parsed.get("missing_criteria") or []
                reply_message = str(parsed.get("reply_message") or "").strip()

                if cleaned_items:
                    return {
                        "items": cleaned_items,
                        "missing_criteria": missing_criteria,
                        "reply_message": reply_message
                    }
            except Exception as e:
                print(f"Groq {model_candidate} error, falling back: {e}")
                continue

    items, missing_criteria = fallback_parse_multi_items(user_msg)
    return {
        "items": items,
        "missing_criteria": missing_criteria,
        "reply_message": ""
    }


@app.route("/api/ai/movement-asker", methods=["POST"])
def ai_movement_asker():
    user = login_required()
    if not user:
        return jsonify({"ok": False, "error": "Not authenticated"}), 401

    if user.role not in {"employee", "manager", "supervisor"}:
        return jsonify({"ok": False, "error": "Unauthorized"}), 403

    data = request.get_json(silent=True) or request.form.to_dict()
    action = data.get("action", "chat")

    # Manual or explicit confirmation branch
    if action == "confirm":
        raw_items = data.get("items")
        if not raw_items or not isinstance(raw_items, list):
            raw_items = [{
                "item": data.get("item"),
                "quantity": data.get("quantity"),
                "direction": data.get("direction", "IN"),
                "note": data.get("note", "")
            }]

        created_movements = []
        for it in raw_items:
            i_name = str(it.get("item", "")).strip().title()
            i_dir = str(it.get("direction", "IN")).strip().upper()
            i_note = str(it.get("note", "")).strip()
            try:
                i_qty = int(it.get("quantity", 0))
            except Exception:
                i_qty = 0

            if i_name and i_qty > 0 and i_dir in {"IN", "OUT"}:
                new_m = Movement(
                    item=i_name, quantity=i_qty, direction=i_dir,
                    note=i_note, created_by=user.user_id
                )
                db.session.add(new_m)
                created_movements.append(new_m)

        if not created_movements:
            return jsonify({"ok": False, "error": "No valid movement criteria to confirm."}), 400

        db.session.commit()
        push_live_movements_to_google_sheet(created_movements)
        trigger_async_sync(delay_seconds=3.0)

        tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
        stats_dict = {
            "total_in": tin,
            "total_out": tout,
            "open_issues": issues_count,
            "employees": employees_count,
            "managers": managers_count
        }

        serialized_logged = [
            {
                "id": m.id,
                "date": m.created_at.strftime("%d %b %Y"),
                "time": m.created_at.strftime("%H:%M:%S"),
                "short_time": m.created_at.strftime("%d %b %H:%M"),
                "item": m.item,
                "quantity": m.quantity,
                "direction": m.direction,
                "status": m.status,
                "created_by": m.created_by,
                "verified_by": m.verified_by or "",
                "note": m.note or ""
            }
            for m in created_movements
        ]

        notify_activity_subscribers({
            "action": "batch_created" if len(serialized_logged) > 1 else "created",
            "movement": serialized_logged[0] if serialized_logged else None,
            "movements": serialized_logged,
            "stats": stats_dict
        })

        return jsonify({
            "ok": True,
            "auto_confirmed": True,
            "message": f"Successfully logged {len(serialized_logged)} movement(s)!",
            "logged_movements": serialized_logged,
            "movement": serialized_logged[0],
            "stats": stats_dict
        })

    # Conversational chat & AUTO-CONFIRMATION by AI
    user_msg = str(data.get("message") or data.get("prompt") or "").strip()
    history = data.get("history", [])
    current_draft = data.get("draft") or {}

    parsed = parse_movement_with_ai(user_msg, history, current_draft)
    items = parsed.get("items", [])
    missing_criteria = parsed.get("missing_criteria", [])

    complete_items = [it for it in items if it.get("is_complete") and it.get("item") and it.get("quantity") and int(it["quantity"]) > 0]
    incomplete_items = [it for it in items if it not in complete_items]

    serialized_logged = []
    tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
    stats_dict = {
        "total_in": tin,
        "total_out": tout,
        "open_issues": issues_count,
        "employees": employees_count,
        "managers": managers_count
    }

    # Automatically add and confirm complete items by AI
    if complete_items:
        created_movements = []
        for it in complete_items:
            new_m = Movement(
                item=it["item"],
                quantity=int(it["quantity"]),
                direction=it["direction"],
                note=it.get("note", ""),
                created_by=user.user_id
            )
            db.session.add(new_m)
            created_movements.append(new_m)

        db.session.commit()
        push_live_movements_to_google_sheet(created_movements)
        trigger_async_sync(delay_seconds=3.0)

        tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
        stats_dict = {
            "total_in": tin,
            "total_out": tout,
            "open_issues": issues_count,
            "employees": employees_count,
            "managers": managers_count
        }

        serialized_logged = [
            {
                "id": m.id,
                "date": m.created_at.strftime("%d %b %Y"),
                "time": m.created_at.strftime("%H:%M:%S"),
                "short_time": m.created_at.strftime("%d %b %H:%M"),
                "item": m.item,
                "quantity": m.quantity,
                "direction": m.direction,
                "status": m.status,
                "created_by": m.created_by,
                "verified_by": m.verified_by or "",
                "note": m.note or ""
            }
            for m in created_movements
        ]

        notify_activity_subscribers({
            "action": "batch_created" if len(serialized_logged) > 1 else "created",
            "movement": serialized_logged[0] if serialized_logged else None,
            "movements": serialized_logged,
            "stats": stats_dict
        })

    # Build crystal-clear response: auto-confirmed items + criteria left
    response_lines = []
    if serialized_logged:
        response_lines.append(f"✅ **Auto-Confirmed & Added {len(serialized_logged)} Movement{'s' if len(serialized_logged) > 1 else ''}:**")
        for m in serialized_logged:
            dir_label = "▲ INCOMING (Received)" if m["direction"] == "IN" else "▼ OUTGOING (Dispatched)"
            note_str = f" • *Notes: {m['note']}*" if m["note"] else ""
            response_lines.append(f"• **{m['quantity']} x {m['item']}** ({dir_label}){note_str}")

        response_lines.append("\n📋 **Criteria Left & Optional Details:**")
        response_lines.append(f"• **Item Name / SKU**: Confirmed ({', '.join(m['item'] for m in serialized_logged)})")
        response_lines.append(f"• **Quantity**: Confirmed ({', '.join(str(m['quantity']) for m in serialized_logged)})")
        response_lines.append(f"• **Movement Direction**: Confirmed ({', '.join(m['direction'] for m in serialized_logged)})")

        unnoted = [m["item"] for m in serialized_logged if not m["note"]]
        if unnoted:
            response_lines.append("• **Notes / Bay / PO (Optional)**: None provided. If you want to attach a Bay, PO, or note, reply here (e.g. *'bay 4'* or *'PO #992'*).")

    if incomplete_items:
        if not serialized_logged:
            response_lines.append("⏳ **Missing Criteria Needed to Auto-Confirm:**")
        else:
            response_lines.append("\n⚠️ **Criteria Left Needed for Incomplete Items:**")

        for inc in incomplete_items:
            name = inc.get("item") or "Item"
            missing = []
            if not inc.get("quantity"):
                missing.append("Quantity")
            if not inc.get("direction"):
                missing.append("Movement Direction")
            if not inc.get("item"):
                missing.append("Item Name / SKU")
            response_lines.append(f"• **{name}**: Missing {', '.join(missing)} (e.g. *'add 50 {name.lower()}'*)")

    if not response_lines:
        response_lines.append("👋 Please specify the items and quantities to log (e.g. *'add 70 sofa, 80 car, 90 wheel'*). I will auto-confirm and add them immediately!")

    final_reply = "\n".join(response_lines)
    latest_item = serialized_logged[-1] if serialized_logged else (incomplete_items[0] if incomplete_items else {})

    return jsonify({
        "ok": True,
        "auto_confirmed": len(serialized_logged) > 0,
        "count": len(serialized_logged),
        "logged_movements": serialized_logged,
        "movement": serialized_logged[-1] if serialized_logged else None,
        "incomplete_items": incomplete_items,
        "missing_criteria": missing_criteria,
        "reply_message": final_reply,
        "stats": stats_dict,
        "latest_draft": {
            "item": latest_item.get("item", ""),
            "quantity": latest_item.get("quantity"),
            "direction": latest_item.get("direction", "IN"),
            "note": latest_item.get("note", "")
        }
    })


@app.route("/movement", methods=["POST"])
def movement():
    user = login_required()
    if not user:
        return redirect(url_for("login"))

    if user.role not in {"employee", "manager", "supervisor"}:
        return role_dashboard(user)

    if "excel_file" in request.files and request.files["excel_file"].filename != "":
        return import_excel_movements()

    if request.is_json:
        req_data = request.get_json(silent=True) or {}
        item = str(req_data.get("item", "")).strip()
        direction = str(req_data.get("direction", "IN")).strip().upper()
        note = str(req_data.get("note", "")).strip()
        try:
            quantity = int(req_data.get("quantity", 0))
        except (ValueError, TypeError):
            quantity = 0
    else:
        item = request.form.get("item", "").strip()
        direction = request.form.get("direction", "IN").upper()
        note = request.form.get("note", "").strip()
        try:
            quantity = int(request.form.get("quantity", "0"))
        except ValueError:
            quantity = 0

    if not item or quantity <= 0 or direction not in {"IN", "OUT"}:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Enter a valid item, quantity and movement type."}), 400
        flash("Enter a valid item, quantity and movement type.", "error")
        return role_dashboard(user)

    new_m = Movement(
        item=item, quantity=quantity, direction=direction,
        note=note, created_by=user.user_id
    )
    db.session.add(new_m)
    db.session.commit()
    push_live_movements_to_google_sheet([new_m])
    trigger_async_sync(delay_seconds=3.0)

    tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
    m_dict = {
        "id": new_m.id,
        "date": new_m.created_at.strftime("%d %b %Y"),
        "time": new_m.created_at.strftime("%H:%M:%S"),
        "short_time": new_m.created_at.strftime("%d %b %H:%M"),
        "item": new_m.item,
        "quantity": new_m.quantity,
        "direction": new_m.direction,
        "status": new_m.status,
        "created_by": new_m.created_by,
        "verified_by": new_m.verified_by or "",
        "note": new_m.note or ""
    }
    stats_dict = {
        "total_in": tin,
        "total_out": tout,
        "open_issues": issues_count,
        "employees": employees_count,
        "managers": managers_count
    }

    notify_activity_subscribers({
        "action": "created",
        "movement": m_dict,
        "stats": stats_dict
    })

    if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
        return jsonify({
            "ok": True,
            "message": f"{quantity} item(s) recorded as {'INCOMING' if direction == 'IN' else 'OUTGOING'}.",
            "movement": m_dict,
            "stats": stats_dict
        })

    flash(f"{quantity} item(s) recorded as {'INCOMING' if direction == 'IN' else 'OUTGOING'}.", "success")
    return role_dashboard(user)


@app.route("/movement/import-excel", methods=["POST"])
def import_excel_movements():
    user = login_required()
    if not user:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Not authenticated"}), 401
        return redirect(url_for("login"))

    if user.role not in {"employee", "manager", "supervisor"}:
        return role_dashboard(user)

    default_direction = request.form.get("default_direction", "IN").strip().upper()
    if default_direction not in {"IN", "OUT"}:
        default_direction = "IN"

    records = []
    errors = []
    file_name = "data"

    extracted_json_str = request.form.get("extracted_json", "")
    if extracted_json_str:
        try:
            records = json.loads(extracted_json_str)
        except Exception:
            records = []

    if not records and "excel_file" in request.files:
        f = request.files["excel_file"]
        if f and f.filename != "":
            file_name = f.filename
            file_bytes = f.read()
            records, errors = extract_movements_from_bytes(file_bytes, file_name, default_direction)

    if not records:
        error_msg = errors[0] if errors else "No valid movement records found in the uploaded file."
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": error_msg, "errors": errors}), 400
        flash(f"Excel Import Failed: {error_msg}", "error")
        return role_dashboard(user)

    created_movements = []
    in_count = 0
    out_count = 0
    for r in records:
        item = str(r.get("item", "")).strip()
        try:
            qty = int(r.get("quantity", 0))
        except (ValueError, TypeError):
            qty = 0
        direction = str(r.get("direction", default_direction)).strip().upper()
        if direction not in {"IN", "OUT"}:
            direction = default_direction
        note = str(r.get("note", "")).strip()

        if item and qty > 0:
            m = Movement(
                item=item,
                quantity=qty,
                direction=direction,
                note=note,
                created_by=user.user_id,
                status="Pending"
            )
            db.session.add(m)
            created_movements.append(m)
            if direction == "IN":
                in_count += 1
            else:
                out_count += 1

    if not created_movements:
        flash("No valid movements could be extracted from the file.", "error")
        return role_dashboard(user)

    db.session.commit()
    push_live_movements_to_google_sheet(created_movements)
    trigger_async_sync(delay_seconds=3.0)

    tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
    stats_dict = {
        "total_in": tin,
        "total_out": tout,
        "open_issues": issues_count,
        "employees": employees_count,
        "managers": managers_count
    }
    serialized_created = [
        {
            "id": m.id,
            "date": m.created_at.strftime("%d %b %Y"),
            "time": m.created_at.strftime("%H:%M:%S"),
            "short_time": m.created_at.strftime("%d %b %H:%M"),
            "item": m.item,
            "quantity": m.quantity,
            "direction": m.direction,
            "status": m.status,
            "created_by": m.created_by,
            "verified_by": m.verified_by or "",
            "note": m.note or ""
        }
        for m in created_movements
    ]
    notify_activity_subscribers({
        "action": "batch_created",
        "movements": serialized_created,
        "stats": stats_dict
    })

    msg = f"Successfully extracted and logged {len(created_movements)} item movements from '{file_name}' ({in_count} IN, {out_count} OUT)!"
    if errors:
        msg += f" Note: {len(errors)} row(s) were skipped due to formatting issues."

    if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
        return jsonify({
            "ok": True,
            "message": msg,
            "count": len(created_movements),
            "movements": serialized_created,
            "stats": stats_dict,
            "in_count": in_count,
            "out_count": out_count,
            "skipped_count": len(errors),
            "errors": errors
        })

    flash(msg, "success")
    return role_dashboard(user)


@app.route("/movement/preview-excel", methods=["POST"])
def preview_excel_movements():
    user = login_required()
    if not user:
        return jsonify({"ok": False, "error": "Not authenticated"}), 401

    if "excel_file" not in request.files:
        return jsonify({"ok": False, "error": "No file uploaded"}), 400

    f = request.files["excel_file"]
    if not f or f.filename == "":
        return jsonify({"ok": False, "error": "No file selected"}), 400

    default_direction = request.form.get("default_direction", "IN").strip().upper()
    if default_direction not in {"IN", "OUT"}:
        default_direction = "IN"

    file_bytes = f.read()
    records, errors = extract_movements_from_bytes(file_bytes, f.filename, default_direction)

    in_count = sum(1 for r in records if r["direction"] == "IN")
    out_count = len(records) - in_count

    return jsonify({
        "ok": True,
        "filename": f.filename,
        "count": len(records),
        "in_count": in_count,
        "out_count": out_count,
        "records": records,
        "errors": errors
    })


@app.route("/movement/sample-template")
def movement_sample_template():
    user = login_required()
    if not user:
        return redirect(url_for("login"))

    fmt = request.args.get("format", "xlsx").lower()
    if fmt == "xlsx":
        import openpyxl, io
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Item Movements"
        ws.append(["Item Name / SKU", "Quantity", "Movement Type", "Additional Notes"])
        ws.append(["Industrial Hydraulic Jack 5-Ton", 10, "IN", "Received from Supplier Alpha"])
        ws.append(["Safety Helmet Pro-Yellow", 45, "OUT", "Dispatched to Sector B"])
        ws.append(["Forklift Battery 48V", 4, "IN", "Restock pallet 12"])
        ws.append(["Storage Bins 20L", 30, "OUT", "Workstation setup"])
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return Response(
            bio.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=sample_item_movements.xlsx"}
        )
    else:
        import io, csv
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Item Name / SKU", "Quantity", "Movement Type", "Additional Notes"])
        writer.writerow(["Industrial Hydraulic Jack 5-Ton", 10, "IN", "Received from Supplier Alpha"])
        writer.writerow(["Safety Helmet Pro-Yellow", 45, "OUT", "Dispatched to Sector B"])
        writer.writerow(["Forklift Battery 48V", 4, "IN", "Restock pallet 12"])
        writer.writerow(["Storage Bins 20L", 30, "OUT", "Workstation setup"])
        return Response(
            output.getvalue(),
            mimetype="text/csv",
            headers={"Content-Disposition": "attachment; filename=sample_item_movements.csv"}
        )


@app.route("/issue", methods=["POST"])
def create_issue():
    user = login_required()
    if not user:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Not authenticated"}), 401
        return redirect(url_for("login"))

    item = request.form.get("item", "").strip()
    issue_type = request.form.get("issue_type", "Other")
    description = request.form.get("description", "").strip()
    photo_file = request.files.get("photo")
    photo_name = ""

    if not item or not description:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Item and problem description are required."}), 400
        flash("Item and problem description are required.", "error")
        return role_dashboard(user)

    if photo_file and photo_file.filename:
        from werkzeug.utils import secure_filename
        import os
        ext = os.path.splitext(photo_file.filename)[1]
        photo_name = secure_filename(f"{uuid.uuid4().hex}{ext}")
        photo_file.save(ISSUES_FOLDER / photo_name)

    new_issue = Issue(
        item=item, issue_type=issue_type, description=description,
        photo=photo_name, created_by=user.user_id, status="Open"
    )
    db.session.add(new_issue)
    db.session.commit()
    trigger_async_sync()

    push_live_event_to_google_sheet("issue", [
        new_issue.id,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Issue",
        new_issue.item,
        new_issue.issue_type,
        "-",
        new_issue.status,
        new_issue.created_by,
        "-",
        new_issue.description or "-"
    ])

    tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
    stats_dict = {
        "total_in": tin,
        "total_out": tout,
        "open_issues": issues_count,
        "employees": employees_count,
        "managers": managers_count
    }
    issue_dict = {
        "id": new_issue.id,
        "item": new_issue.item,
        "issue_type": new_issue.issue_type,
        "description": new_issue.description,
        "photo": new_issue.photo or "",
        "created_by": new_issue.created_by,
        "status": "Open",
        "created_at": new_issue.created_at.strftime("%d %b %Y"),
        "comments": []
    }

    notify_activity_subscribers({
        "action": "issue_created",
        "issue": issue_dict,
        "stats": stats_dict
    })

    if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
        return jsonify({
            "ok": True,
            "message": "Alert sent to Management.",
            "issue": issue_dict,
            "stats": stats_dict
        })

    flash("Alert sent to Management.", "success")
    return role_dashboard(user)


@app.route("/issue/<int:issue_id>/comment", methods=["POST"])
def add_issue_comment(issue_id):
    user = login_required()
    if not user:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Not authenticated"}), 401
        return redirect(url_for("login"))
        
    if user.role == "supervisor":
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Supervisors can only view conversations, not participate."}), 403
        flash("Supervisors can only view conversations, not participate.", "error")
        return role_dashboard(user)
        
    body = request.form.get("body", "").strip()
    if not body:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Comment cannot be empty."}), 400
        return role_dashboard(user)

    comment = IssueComment(issue_id=issue_id, sender_id=user.user_id, body=body)
    db.session.add(comment)
    db.session.commit()
    trigger_async_sync()

    comment_dict = {
        "id": comment.id,
        "issue_id": issue_id,
        "sender_id": comment.sender_id,
        "body": comment.body,
        "created_at": comment.created_at.strftime("%d %b %H:%M"),
        "time": comment.created_at.strftime("%H:%M")
    }

    notify_activity_subscribers({
        "action": "issue_commented",
        "issue_id": issue_id,
        "comment": comment_dict
    })

    if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
        return jsonify({"ok": True, "comment": comment_dict})

    flash("Comment added to issue.", "success")
    return role_dashboard(user)


@app.route("/issue/<int:issue_id>/resolve", methods=["POST"])
def resolve_issue(issue_id):
    user = login_required()
    if not user:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Not authenticated"}), 401
        return redirect(url_for("login"))
    if user.role != "manager":
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Only Managers have the authority to resolve issues."}), 403
        flash("Only Managers have the authority to resolve issues.", "error")
        return role_dashboard(user)

    issue = db.session.get(Issue, issue_id)
    if issue:
        issue.status = "Resolved"
        issue.resolved_by = user.user_id
        
        supervisors = User.query.filter_by(role="supervisor").all()
        for sup in supervisors:
            msg = Message(
                sender_id="SYSTEM",
                receiver_id=sup.user_id,
                subject=f"Issue Resolved: {issue.item}",
                body=f"Manager {user.user_id} has resolved the {issue.issue_type} issue reported by {issue.created_by}. The full conversation is available in your Resolved Issues Archive.",
                msg_type="confidential"
            )
            db.session.add(msg)
            
        db.session.commit()
        trigger_async_sync()

        tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
        stats_dict = {
            "total_in": tin,
            "total_out": tout,
            "open_issues": issues_count,
            "employees": employees_count,
            "managers": managers_count
        }

        notify_activity_subscribers({
            "action": "issue_resolved",
            "issue_id": issue.id,
            "item": issue.item,
            "issue_type": issue.issue_type,
            "description": issue.description,
            "photo": issue.photo or "",
            "created_by": issue.created_by,
            "resolved_by": user.user_id,
            "date": issue.created_at.strftime("%d %b %Y"),
            "comments": [
                {
                    "sender_id": c.sender_id,
                    "body": c.body,
                    "created_at": c.created_at.strftime("%d %b %H:%M")
                }
                for c in issue.comments
            ],
            "stats": stats_dict
        })

        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({
                "ok": True,
                "message": "Issue resolved.",
                "issue_id": issue.id,
                "resolved_by": user.user_id,
                "stats": stats_dict
            })

        flash("Issue resolved.", "success")
    return role_dashboard(user)


@app.route("/message/send", methods=["POST"])
def send_message():
    user = login_required()
    if not user:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Not authenticated"}), 401
        return redirect(url_for("login"))

    msg_type = request.form.get("msg_type", "confidential").lower()
    receiver_id = request.form.get("receiver_id", "").strip().upper()
    subject = request.form.get("subject", "").strip()
    if not subject:
        subject = "Direct Message"
    body = request.form.get("body", "").strip()

    if not body:
        if is_ajax_request():
            return jsonify({"ok": False, "error": "Message body is required."}), 400
        flash("Message body is required.", "error")
        return role_dashboard(user)

    if msg_type == "common":
        if user.role not in {"supervisor", "manager"}:
            if is_ajax_request():
                return jsonify({"ok": False, "error": "Only Managers and Supervisors can broadcast common messages."}), 403
            flash("Only Managers and Supervisors can broadcast common messages.", "error")
            return role_dashboard(user)

        new_msg = Message(
            sender_id=user.user_id,
            receiver_id="ALL",
            subject=subject,
            body=body,
            is_read=False,
            msg_type="common"
        )
        db.session.add(new_msg)
        db.session.commit()

        msg_dict = {
            "id": new_msg.id,
            "sender_id": new_msg.sender_id,
            "receiver_id": "ALL",
            "subject": new_msg.subject,
            "body": new_msg.body,
            "created_at": new_msg.created_at.strftime("%d %b %H:%M"),
            "msg_type": "common"
        }

        notify_activity_subscribers({
            "action": "common_broadcast",
            "message": msg_dict
        })

        if is_ajax_request():
            return jsonify({
                "ok": True,
                "message": "Common message broadcasted successfully.",
                "broadcast": msg_dict
            })

        flash("Common message broadcasted successfully.", "success")
        return role_dashboard(user)

    receiver = User.query.filter_by(user_id=receiver_id, active=True).first()
    if not receiver:
        if is_ajax_request():
            return jsonify({"ok": False, "error": "Receiver not found or inactive."}), 404
        flash("Receiver not found or inactive.", "error")
        return role_dashboard(user)

    if receiver.user_id == user.user_id:
        if is_ajax_request():
            return jsonify({"ok": False, "error": "You cannot send a message to yourself."}), 400
        flash("You cannot send a message to yourself.", "error")
        return role_dashboard(user)

    new_msg = Message(
        sender_id=user.user_id,
        receiver_id=receiver.user_id,
        subject=subject,
        body=body,
        is_read=False,
        msg_type="confidential"
    )
    db.session.add(new_msg)
    db.session.commit()

    unread_count = Message.query.filter_by(receiver_id=receiver.user_id, msg_type="confidential", is_read=False).count()

    msg_dict = {
        "id": new_msg.id,
        "sender_id": new_msg.sender_id,
        "receiver_id": new_msg.receiver_id,
        "receiver_name": receiver.name,
        "subject": new_msg.subject,
        "body": new_msg.body,
        "attachment": new_msg.attachment or "",
        "created_at": new_msg.created_at.strftime("%d %b %H:%M"),
        "is_read": False,
        "msg_type": "confidential"
    }

    notify_activity_subscribers({
        "action": "confidential_message",
        "message": msg_dict,
        "receiver_unread": unread_count
    })

    if is_ajax_request():
        return jsonify({
            "ok": True,
            "message": f"Confidential message sent to {receiver.name}.",
            "message_data": msg_dict
        })

    flash(f"Confidential message sent to {receiver.name}.", "success")
    return role_dashboard(user)

@app.route("/movement/<int:movement_id>/verify", methods=["POST"])
def verify_movement(movement_id):
    user = login_required()
    if not user or user.role != "manager":
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Only Managers can verify movements."}), 403
        flash("Only Managers can verify movements.", "error")
        return redirect(url_for("login"))
    
    movement = db.session.get(Movement, movement_id)
    if movement:
        action = request.form.get("action", "Verified")
        movement.status = action
        movement.verified_by = user.user_id
        db.session.commit()
        update_google_sheet_movement_status(
            movement.id,
            action,
            user.user_id,
            f"Action: {action} by {user.user_id}",
            fallback_row=[
                movement.id,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "Verification",
                movement.item,
                movement.quantity,
                movement.direction,
                action,
                movement.created_by,
                user.user_id,
                f"Action: {action} by {user.user_id}"
            ]
        )
        trigger_async_sync(delay_seconds=3.0)
        tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
        stats_dict = {
            "total_in": tin,
            "total_out": tout,
            "open_issues": issues_count,
            "employees": employees_count,
            "managers": managers_count
        }
        notify_activity_subscribers({
            "action": "verified",
            "id": movement.id,
            "status": action,
            "verified_by": user.user_id,
            "stats": stats_dict
        })
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({
                "ok": True,
                "message": f"Movement #{movement.id} marked as {action}.",
                "id": movement.id,
                "status": action,
                "verified_by": user.user_id,
                "stats": stats_dict
            })
        flash(f"Movement marked as {action}.", "success")
    return role_dashboard(user)


@app.route("/google_sheet/sync", methods=["POST"])
def google_sheet_sync():
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    if user.role not in {"supervisor", "manager"}:
        flash("Only Supervisors and Managers can sync with Google Sheets.", "error")
        return role_dashboard(user)

    webhook_url = request.form.get("webhook_url", "").strip()
    if webhook_url:
        conf = get_sheet_config()
        conf["webhook_url"] = webhook_url
        save_sheet_config(conf)

    success, msg = sync_all_to_google_sheet(webhook_url)
    if success:
        flash(msg, "success")
    else:
        flash(msg, "error")
    return role_dashboard(user)


@app.route("/api/sheets/live_activity.csv")
def live_activity_csv():
    import csv, io
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["ID", "Timestamp", "Activity", "Item", "Quantity", "Direction", "Status", "Logged By", "Verified By", "Notes"])
    movements = Movement.query.order_by(Movement.created_at.desc()).all()
    for m in movements:
        writer.writerow([m.id, m.created_at.strftime("%Y-%m-%d %H:%M:%S"), "Movement", m.item, m.quantity, m.direction, m.status, m.created_by, m.verified_by or "-", m.note or "-"])
    
    issues = Issue.query.order_by(Issue.created_at.desc()).all()
    for i in issues:
        writer.writerow([i.id, i.created_at.strftime("%Y-%m-%d %H:%M:%S"), "Issue", i.item, 1, "-", i.status, i.created_by, i.resolved_by or "-", i.description or "-"])
        
    return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "inline; filename=warehouse_live_activity.csv"})


@app.route("/api/sheets/manager_report.csv")
def manager_report_csv():
    import csv, io
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Manager Name", "User ID", "Work Status", "Account Status", "Logs Verified", "Issues Resolved", "Performance Score", "Last Active Timestamp"])
    manager_users = User.query.filter_by(role="manager").order_by(User.name).all()
    for mgr in manager_users:
        verifications = Movement.query.filter_by(verified_by=mgr.user_id).count()
        resolutions = Issue.query.filter_by(resolved_by=mgr.user_id).count()
        score = (verifications * 2) + (resolutions * 5)
        work_status = "Online (Working)" if mgr.is_online else "Logged Out"
        account_status = "Active" if mgr.active else "Disabled"
        last_seen = mgr.last_seen.strftime("%Y-%m-%d %H:%M:%S") if mgr.last_seen else "Never"
        writer.writerow([mgr.name, mgr.user_id, work_status, account_status, verifications, resolutions, score, last_seen])
    return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "inline; filename=manager_report.csv"})


@app.route("/api/sheets/employee_report.csv")
def employee_report_csv():
    import csv, io
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Employee Name", "Employee ID", "Work Status", "Action Date & Time", "Item Logged", "Quantity", "Direction", "Verification Status", "Verified By", "Notes"])
    employee_users = User.query.filter_by(role="employee").order_by(User.name).all()
    for emp in employee_users:
        emp_movements = Movement.query.filter_by(created_by=emp.user_id).order_by(Movement.created_at.desc()).all()
        work_status = "Online (Working)" if emp.is_online else "Logged Out"
        if emp_movements:
            for m in emp_movements:
                writer.writerow([emp.name, emp.user_id, work_status, m.created_at.strftime("%Y-%m-%d %H:%M:%S"), m.item, m.quantity, m.direction, m.status, m.verified_by or "Pending Verification", m.note or "-"])
        else:
            writer.writerow([emp.name, emp.user_id, work_status, "-", "No Movements Logged Yet", 0, "-", "-", "-", "-"])
    return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "inline; filename=employee_report.csv"})


@app.route("/api/sheets/manager_performance.csv")
def manager_performance_csv():
    import csv, io
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Manager Name", "Manager ID", "Work Status", "Account Status", "Logs Verified", "Issues Resolved", "Flags Raised", "Performance Score", "Last Active Timestamp"])
    records = get_manager_performance_data()
    for r in records:
        writer.writerow([r["name"], r["user_id"], r["work_status"], r["account_status"], r["verifications"], r["resolutions"], r["flags_raised"], r["score"], r["last_seen"]])
    return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "inline; filename=manager_performance.csv"})


@app.route("/api/sheets/employee_performance.csv")
def employee_performance_csv():
    import csv, io
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Employee Name", "Employee ID", "Work Status", "Account Status", "Total Logs", "Verified Logs", "Flagged Logs", "Pending Logs", "Total Units Moved", "Accuracy Rate", "Performance Score", "Last Active Timestamp"])
    records = get_employee_performance_data()
    for r in records:
        writer.writerow([r["name"], r["user_id"], r["work_status"], r["account_status"], r["total_logs"], r["verified_logs"], r["flagged_logs"], r["pending_logs"], r["total_units"], r["accuracy_rate"], r["score"], r["last_seen"]])
    return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "inline; filename=employee_performance.csv"})


@app.route("/export/performance/<role>", methods=["GET", "POST"])
def export_performance(role):
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    if role not in ["manager", "employee"]:
        flash("Invalid performance role.", "error")
        return redirect(url_for("dashboard"))
    if role == "manager" and user.role not in ["supervisor", "manager"]:
        flash("Unauthorized to export manager performance.", "error")
        return redirect(url_for("dashboard"))
    if user.role not in ["supervisor", "manager"]:
        flash("Unauthorized.", "error")
        return redirect(url_for("dashboard"))

    if role == "manager":
        res = manager_performance_csv()
        res.headers["Content-Disposition"] = f"attachment; filename=manager_performance_{int(datetime.now().timestamp())}.csv"
        return res
    else:
        res = employee_performance_csv()
        res.headers["Content-Disposition"] = f"attachment; filename=employee_performance_{int(datetime.now().timestamp())}.csv"
        return res


@app.route("/movement/<int:movement_id>/delete", methods=["POST"])
def delete_movement(movement_id):
    user = login_required()
    if not user:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Not authenticated."}), 401
        return redirect(url_for("login"))
    
    movement = db.session.get(Movement, movement_id)
    if not movement:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Movement record not found."}), 404
        flash("Movement record not found.", "error")
        return role_dashboard(user)

    if user.role == "employee" and movement.created_by != user.user_id:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "You can only delete movements you logged."}), 403
        flash("You can only delete movements you logged.", "error")
        return redirect(url_for("employee_dashboard"))

    if user.role not in {"manager", "supervisor", "employee"}:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Unauthorized action."}), 403
        flash("Unauthorized action.", "error")
        return redirect(url_for("login"))
    
    if movement:
        movement.status = "Deleted"
        movement.verified_by = user.user_id
        db.session.commit()
        update_google_sheet_movement_status(
            movement.id,
            "Deleted",
            user.user_id,
            f"Deleted by {user.user_id}",
            fallback_row=[
                movement.id,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "Deleted",
                movement.item,
                movement.quantity,
                movement.direction,
                "Deleted",
                movement.created_by,
                user.user_id,
                f"Deleted by {user.user_id}"
            ]
        )
        trigger_async_sync(delay_seconds=3.0)
        tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
        stats_dict = {
            "total_in": tin,
            "total_out": tout,
            "open_issues": issues_count,
            "employees": employees_count,
            "managers": managers_count
        }
        notify_activity_subscribers({
            "action": "manager_deleted",
            "id": movement_id,
            "item": movement.item,
            "quantity": movement.quantity,
            "direction": movement.direction,
            "status": "Deleted",
            "deleted_by": user.user_id,
            "verified_by": user.user_id,
            "date": movement.created_at.strftime("%d %b %Y"),
            "time": movement.created_at.strftime("%H:%M:%S"),
            "created_by": movement.created_by,
            "note": movement.note or "",
            "stats": stats_dict
        })
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({
                "ok": True,
                "message": "Movement record marked as deleted.",
                "id": movement_id,
                "item": movement.item,
                "quantity": movement.quantity,
                "direction": movement.direction,
                "status": "Deleted",
                "deleted_by": user.user_id,
                "verified_by": user.user_id,
                "stats": stats_dict
            })
        flash("Movement record deleted.", "success")
    return role_dashboard(user)

@app.route("/movement/<int:movement_id>/hard_delete", methods=["POST"])
def hard_delete_movement(movement_id):
    user = login_required()
    if not user or user.role != "supervisor":
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Unauthorized action."}), 403
        flash("Unauthorized action.", "error")
        return redirect(url_for("login"))
    
    movement = db.session.get(Movement, movement_id)
    if movement:
        db.session.delete(movement)
        db.session.commit()
        update_google_sheet_movement_status(
            movement_id,
            "Hard Deleted",
            user.user_id,
            f"Permanently removed by {user.user_id}",
            fallback_row=[
                movement_id,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "Hard Deleted",
                "-",
                0,
                "-",
                "Deleted",
                "-",
                user.user_id,
                f"Permanently removed by {user.user_id}"
            ]
        )
        trigger_async_sync(delay_seconds=3.0)
        tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
        stats_dict = {
            "total_in": tin,
            "total_out": tout,
            "open_issues": issues_count,
            "employees": employees_count,
            "managers": managers_count
        }
        notify_activity_subscribers({
            "action": "hard_deleted",
            "id": movement_id,
            "stats": stats_dict
        })
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": True, "message": "Movement permanently deleted.", "id": movement_id, "stats": stats_dict})
        flash("Movement permanently deleted.", "success")
    return role_dashboard(user)

@app.route("/movement/<int:movement_id>/restore", methods=["POST"])
def restore_movement(movement_id):
    user = login_required()
    if not user or user.role != "supervisor":
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Unauthorized action. Only Supervisors can restore movements."}), 403
        flash("Unauthorized action.", "error")
        return redirect(url_for("login"))
    
    movement = db.session.get(Movement, movement_id)
    if movement:
        movement.status = "Pending"
        movement.verified_by = None
        db.session.commit()
        update_google_sheet_movement_status(
            movement.id,
            "Pending",
            "",
            f"Restored by {user.user_id}",
            fallback_row=[
                movement.id,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "Restored",
                movement.item,
                movement.quantity,
                movement.direction,
                "Pending",
                movement.created_by,
                user.user_id,
                f"Restored by {user.user_id}"
            ]
        )
        trigger_async_sync(delay_seconds=3.0)
        tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
        stats_dict = {
            "total_in": tin,
            "total_out": tout,
            "open_issues": issues_count,
            "employees": employees_count,
            "managers": managers_count
        }
        mov_dict = {
            "id": movement.id,
            "date": movement.created_at.strftime("%d %b %Y"),
            "time": movement.created_at.strftime("%H:%M:%S"),
            "item": movement.item,
            "quantity": movement.quantity,
            "direction": movement.direction,
            "status": movement.status,
            "created_by": movement.created_by,
            "verified_by": "",
            "note": movement.note or ""
        }
        notify_activity_subscribers({
            "action": "restored",
            "id": movement.id,
            "movement": mov_dict,
            "stats": stats_dict
        })
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": True, "message": "Movement restored successfully.", "movement": mov_dict, "stats": stats_dict})
        flash("Movement restored.", "success")
    return role_dashboard(user)

@app.route("/issue/<int:issue_id>/delete", methods=["POST"])
def delete_issue(issue_id):
    user = login_required()
    if not user or user.role != "supervisor":
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Unauthorized action."}), 403
        flash("Unauthorized action.", "error")
        return redirect(url_for("login"))
    
    issue = db.session.get(Issue, issue_id)
    if issue:
        IssueComment.query.filter_by(issue_id=issue.id).delete()
        db.session.delete(issue)
        db.session.commit()
        trigger_async_sync()

        tin, tout, issues_count, employees_count, managers_count = dashboard_stats()
        stats_dict = {
            "total_in": tin,
            "total_out": tout,
            "open_issues": issues_count,
            "employees": employees_count,
            "managers": managers_count
        }

        notify_activity_subscribers({
            "action": "issue_deleted",
            "issue_id": issue_id,
            "stats": stats_dict
        })

        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({
                "ok": True,
                "message": "Issue permanently deleted.",
                "issue_id": issue_id,
                "stats": stats_dict
            })

        flash("Issue permanently deleted.", "success")
    return role_dashboard(user)


@app.route("/messages/read/<int:message_id>", methods=["POST"])
def read_message(message_id):
    user = login_required()
    if not user:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Not authenticated"}), 401
        return redirect(url_for("login"))

    msg = db.session.get(Message, message_id)
    if msg and msg.receiver_id == user.user_id:
        msg.is_read = True
        db.session.commit()
        unread_count = Message.query.filter_by(receiver_id=user.user_id, msg_type="confidential", is_read=False).count()

        notify_activity_subscribers({
            "action": "message_read",
            "id": message_id,
            "receiver_id": user.user_id,
            "unread": unread_count
        })

        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({
                "ok": True,
                "id": message_id,
                "unread": unread_count
            })

    return role_dashboard(user)

@app.route("/conversation/<u1>/<u2>")
def view_conversation(u1, u2):
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    
    # Security check: Only supervisors or the actual participants can view
    if user.role != "supervisor" and user.user_id not in [u1, u2]:
        flash("Unauthorized to view this conversation.", "error")
        return role_dashboard(user)
    
    messages = Message.query.filter(
        db.or_(
            db.and_(Message.sender_id == u1, Message.receiver_id == u2),
            db.and_(Message.sender_id == u2, Message.receiver_id == u1)
        ),
        Message.msg_type == "confidential",
        Message.attachment == ""
    ).order_by(Message.created_at.asc()).all()
    
    if user.role != "supervisor":
        messages = [m for m in messages if not (
            (m.sender_id == user.user_id and m.deleted_by_sender) or 
            (m.receiver_id == user.user_id and m.deleted_by_receiver)
        )]
    
    return render_template("conversation.html", user=user, u1=u1, u2=u2, messages=messages)

@app.route("/conversation/delete/<u1>/<u2>", methods=["POST"])
def delete_conversation(u1, u2):
    user = login_required()
    if not user:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Not authenticated"}), 401
        return redirect(url_for("login"))
        
    if user.role != "supervisor" and user.user_id not in [u1, u2]:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Unauthorized to delete this conversation."}), 403
        flash("Unauthorized to delete this conversation.", "error")
        return role_dashboard(user)
        
    messages = Message.query.filter(
        db.or_(
            db.and_(Message.sender_id == u1, Message.receiver_id == u2),
            db.and_(Message.sender_id == u2, Message.receiver_id == u1)
        ),
        Message.msg_type == "confidential"
    ).all()
    
    for msg in messages:
        if user.role == "supervisor":
            db.session.delete(msg)
        else:
            if msg.sender_id == user.user_id:
                msg.deleted_by_sender = True
            if msg.receiver_id == user.user_id:
                msg.deleted_by_receiver = True
        
    db.session.commit()

    notify_activity_subscribers({
        "action": "conversation_deleted",
        "u1": u1,
        "u2": u2
    })

    if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
        return jsonify({
            "ok": True,
            "message": "Conversation deleted successfully.",
            "u1": u1,
            "u2": u2
        })

    flash("Conversation deleted successfully.", "success")
    return role_dashboard(user)

@app.route("/messages/delete/<int:message_id>", methods=["POST"])
def delete_message(message_id):
    user = login_required()
    if not user:
        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({"ok": False, "error": "Not authenticated"}), 401
        return redirect(url_for("login"))

    message = db.session.get(Message, message_id)
    if message:
        msg_type = message.msg_type
        sender_id = message.sender_id
        receiver_id = message.receiver_id

        if user.role == "supervisor":
            db.session.delete(message)
        elif message.sender_id == user.user_id or message.receiver_id == user.user_id:
            if message.sender_id == user.user_id:
                message.deleted_by_sender = True
            if message.receiver_id == user.user_id:
                message.deleted_by_receiver = True
        else:
            if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
                return jsonify({"ok": False, "error": "Unauthorized to delete this message."}), 403
            flash("Unauthorized to delete this message.", "error")
            return role_dashboard(user)

        db.session.commit()

        notify_activity_subscribers({
            "action": "message_deleted",
            "id": message_id,
            "msg_type": msg_type,
            "sender_id": sender_id,
            "receiver_id": receiver_id
        })

        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json:
            return jsonify({
                "ok": True,
                "message": "Message deleted.",
                "id": message_id,
                "msg_type": msg_type
            })

        flash("Message deleted.", "success")
            
    return role_dashboard(user)


@app.route("/users")
def users():
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    if user.role not in {"supervisor", "manager"}:
        flash("Unauthorized.", "error")
        return role_dashboard(user)

    if user.role == "manager":
        user_list = User.query.filter_by(role="employee", active=True).order_by(User.name).all()
    else:
        user_list = User.query.order_by(User.role, User.name).all()

    active_count = sum(1 for u in user_list if u.is_online)
    offline_count = len(user_list) - active_count
    enabled_count = sum(1 for u in user_list if u.active)
    disabled_count = len(user_list) - enabled_count
    sup_count = sum(1 for u in user_list if u.role == 'supervisor')
    mgr_count = sum(1 for u in user_list if u.role == 'manager')
    emp_count = sum(1 for u in user_list if u.role == 'employee')

    return render_template(
        "users.html",
        user=user,
        users=user_list,
        active_count=active_count,
        offline_count=offline_count,
        enabled_count=enabled_count,
        disabled_count=disabled_count,
        sup_count=sup_count,
        mgr_count=mgr_count,
        emp_count=emp_count,
        sheet_config=get_sheet_config()
    )


@app.route("/users/create", methods=["GET", "POST"])
def create_user():
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    if user.role != "supervisor":
        flash("Only Supervisor can create users.", "error")
        return role_dashboard(user)

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        uid = request.form.get("user_id", "").strip().upper()
        password = request.form.get("password", "")
        role = request.form.get("role", "employee")

        if not name or not uid or not password:
            flash("Name, User ID and password are required.", "error")
            return redirect(url_for("users") if request.headers.get("Referer") and "/users" in request.headers.get("Referer") else url_for("create_user"))

        if role not in {"employee", "manager", "supervisor"}:
            flash("Invalid role.", "error")
            return redirect(url_for("users") if request.headers.get("Referer") and "/users" in request.headers.get("Referer") else url_for("create_user"))

        if User.query.filter_by(user_id=uid).first():
            flash("User ID already exists.", "error")
            return redirect(url_for("users") if request.headers.get("Referer") and "/users" in request.headers.get("Referer") else url_for("create_user"))

        new_user = User(
            user_id=uid, name=name,
            password=generate_password_hash(password),
            role=role, qr_token=str(uuid.uuid4()), active=True
        )
        db.session.add(new_user)
        db.session.commit()
        trigger_async_sync()
        generate_qr(new_user)
        flash(f"{name} created and QR generated.", "success")
        return redirect(url_for("users"))

    return render_template("create_user.html", user=user)


@app.route("/users/<int:user_id>/edit", methods=["GET", "POST"])
def edit_user(user_id):
    current_user = login_required()
    if not current_user:
        return redirect(url_for("login"))

    target = db.session.get(User, user_id)
    if not target:
        flash("User not found.", "error")
        return role_dashboard(current_user)

    # Supervisor can edit anyone; Manager can edit employees or themselves; Users can edit themselves
    allowed = False
    if current_user.role == "supervisor":
        allowed = True
    elif current_user.role == "manager" and (target.role == "employee" or target.id == current_user.id):
        allowed = True
    elif target.id == current_user.id:
        allowed = True

    if not allowed:
        flash("Unauthorized to edit this user.", "error")
        return role_dashboard(current_user)

    if request.method == "POST":
        new_name = request.form.get("name", "").strip()
        new_password = request.form.get("password", "").strip()
        new_role = request.form.get("role", "").strip()

        if not new_name:
            flash("User name cannot be empty.", "error")
            return redirect(request.referrer or url_for("users"))

        target.name = new_name
        if new_password:
            target.password = generate_password_hash(new_password)

        if current_user.role == "supervisor" and new_role in {"employee", "manager", "supervisor"}:
            if target.user_id != "SUP001":
                target.role = new_role

        db.session.commit()
        trigger_async_sync()

        if current_user.id == target.id:
            session["name"] = target.name

        msg = f"User '{target.name}' ({target.user_id}) updated successfully."
        flash(msg, "success")

        if request.headers.get("X-Requested-With") == "XMLHttpRequest" or "application/json" in request.headers.get("Accept", ""):
            return jsonify({
                "success": True,
                "message": msg,
                "user": {"id": target.id, "user_id": target.user_id, "name": target.name, "role": target.role}
            })

        ref = request.referrer or ""
        if "/edit" in ref:
            return redirect(url_for("users"))
        return redirect(ref or url_for("users"))

    return render_template("edit_user.html", user=current_user, target=target)


@app.route("/users/<int:user_id>/delete", methods=["POST"])
def delete_user(user_id):
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    if user.role != "supervisor":
        flash("Only Supervisor can delete users.", "error")
        return role_dashboard(user)

    target = db.session.get(User, user_id)
    if not target or target.user_id == "SUP001" or target.id == user.id:
        flash("The main supervisor cannot be deleted.", "error")
        return redirect(url_for("users"))

    qr = qr_file(target)
    db.session.delete(target)
    db.session.commit()
    trigger_async_sync()
    if qr.exists():
        qr.unlink()

    flash(f"{target.name} deleted.", "success")
    return redirect(url_for("users"))


@app.route("/users/<int:user_id>/toggle", methods=["POST"])
def toggle_user(user_id):
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    if user.role != "supervisor":
        flash("Only Supervisor can change user status.", "error")
        return role_dashboard(user)

    target = db.session.get(User, user_id)
    if not target or target.user_id == "SUP001":
        flash("Main supervisor cannot be disabled.", "error")
        return redirect(url_for("users"))

    target.active = not target.active
    db.session.commit()
    trigger_async_sync()
    return redirect(url_for("users"))


@app.route("/users/<int:user_id>/qr")
def regenerate_qr(user_id):
    user = login_required()
    if not user:
        return redirect(url_for("login"))
    if user.role != "supervisor":
        flash("Only Supervisor can regenerate QR.", "error")
        return role_dashboard(user)

    target = db.session.get(User, user_id)
    if not target:
        flash("User not found.", "error")
        return redirect(url_for("users"))

    target.qr_token = str(uuid.uuid4())
    db.session.commit()
    generate_qr(target)
    flash("QR regenerated.", "success")
    return redirect(url_for("users"))


@app.route("/qr/<token>")
def qr_profile(token):
    user = User.query.filter_by(qr_token=token).first()
    if not user:
        return "<h1>Invalid QR Code</h1>", 404
    return render_template("qr_profile.html", user=user)


@app.route("/google_sheets_sync_script.js")
def get_apps_script_js():
    js_path = BASE_DIR / "google_sheets_sync_script.js"
    if js_path.exists():
        with open(js_path, "r", encoding="utf-8") as f:
            return Response(f.read(), mimetype="text/javascript")
    return "File not found", 404


init_db()

try:
    from migrate_db import run_migrations
    run_migrations()
except ImportError:
    print("No migrate_db script found or it could not be imported.")

# Start automated background synchronization daemon
start_google_sheet_autosync()

if __name__ == "__main__":
    print("=" * 60)
    print("WAREHOUSE MANAGEMENT")
    print("Supervisor: SUP001 / admin123")
    print("Open: http://127.0.0.1:5000")
    print("=" * 60)
    app.run(host="0.0.0.0", port=5000, debug=True)
